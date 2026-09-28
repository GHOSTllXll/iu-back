# backend/ai_service/document_classifier.py
"""
Intelligent Document Classification Layer.

Sits at the very front of the CRE ingestion pipeline, before a file is ever
assigned to the Offering Memorandum / Trailing-12 / Rent Roll slot. The
frontend's single "drop your documents here" zone (dashboard/index.vue) calls
this once per dropped file, via ClassifyDocumentView, and uses the result to
auto-place the file in the right slot instead of forcing the user to know
which box a given PDF belongs in — the goal is eliminating the upload
mistakes that come from a user misjudging a document by its filename alone,
and supporting a mixed portfolio data room (e.g. a debt fund's data room)
where files aren't pre-sorted at all.

This is a fast, local keyword scan over a SMALL sample of the document (the
first couple of PDF pages, a handful of Word paragraphs, or the header row +
a few sample rows of a spreadsheet) — never a full-document parse and never
an AI call. It only decides where a file is routed; the actual OM/T12/Rent
Roll extraction and analysis (process_underwriting_files in views.py) is
completely unchanged and runs exactly as it did before this layer existed.

DESIGN PRINCIPLE — fail loud, not silent (same principle as
ppm_splitter.py's page-locating heuristic): this scan has not been tuned
against a broad sample of real-world sponsor/lender templates. When it can't
tell two document types apart, or finds too weak a signal to be confident,
it returns UNCLASSIFIED rather than guessing — a wrong guess here would
silently feed a Rent Roll into the T12 slot (or vice versa) and corrupt the
underwriting numbers downstream, which is a worse outcome than asking the
user to place one file manually.
"""
import os

DOCUMENT_TYPE_RENT_ROLL = 'RENT_ROLL'
DOCUMENT_TYPE_TRAILING_12 = 'TRAILING_12'
DOCUMENT_TYPE_OFFERING_MEMORANDUM = 'OFFERING_MEMORANDUM'
DOCUMENT_TYPE_UNCLASSIFIED = 'UNCLASSIFIED'

# Union of every extension any of the 3 slots accepts today (see
# parsers.ALLOWED_DOC_EXTENSIONS / ALLOWED_EXCEL_EXTENSIONS) — a file outside
# this set can't belong to any slot, so it's rejected before classification
# is even attempted.
CLASSIFIABLE_EXTENSIONS = {'.pdf', '.doc', '.docx', '.xls', '.xlsx', '.csv'}

# How many PDF pages / Word paragraphs' worth of content to sample.
# Deliberately small — classification only needs a title page or a table's
# header row, not the whole document, and this runs once per dropped file so
# it needs to stay fast.
PDF_PAGES_TO_SCAN = 2
DOCX_PARAGRAPHS_TO_SCAN = 40
EXCEL_SAMPLE_ROWS = 10

# Each entry in a rule is a group of interchangeable phrases — a hit on ANY
# phrase in a group counts as one signal for that group. Scoring counts
# DISTINCT GROUPS matched, not raw keyword occurrences, so a document can't
# be tagged off one phrase repeated many times; it has to show several of the
# structural signals a real document of that type would carry.
#
# The first phrase in each group is the literal keyword from the original
# spec; the rest are common real-world variants added so the heuristic isn't
# defeated by a sponsor using a slightly different label for the same column.
RENT_ROLL_KEYWORD_GROUPS = [
    ['unit', 'unit #', 'unit no.', 'suite'],
    ['tenant name', 'tenant', 'resident name', 'resident'],
    ['lease start', 'lease begin', 'move-in', 'move in date', 'lease commencement'],
    ['monthly rent', 'rent amount', 'current rent', 'base rent'],
    ['arrears', 'balance due', 'past due', 'amount owed'],
]

TRAILING_12_KEYWORD_GROUPS = [
    ['revenue', 'income'],
    ['jan', 'january'],
    ['feb', 'february'],
    ['total ytd', 'year to date', 'ytd total', 'ytd'],
    ['net operating income', 'noi'],
]

OFFERING_MEMORANDUM_KEYWORD_GROUPS = [
    ['offering memorandum', 'confidential offering', 'confidential memorandum', 'investment offering'],
    ['executive summary', 'investment summary', 'investment highlights'],
    ['property description', 'property overview', 'the offering'],
    ['asking price', 'purchase price', 'list price', 'offering price'],
]

CLASSIFICATION_RULES = [
    (DOCUMENT_TYPE_RENT_ROLL, RENT_ROLL_KEYWORD_GROUPS),
    (DOCUMENT_TYPE_TRAILING_12, TRAILING_12_KEYWORD_GROUPS),
    (DOCUMENT_TYPE_OFFERING_MEMORANDUM, OFFERING_MEMORANDUM_KEYWORD_GROUPS),
]

# A document must match at least this many DISTINCT groups to be considered
# at all, and its top score must beat the runner-up outright (no ties) —
# otherwise the result is UNCLASSIFIED. See module docstring's fail-loud note.
MIN_DISTINCT_GROUP_HITS = 2

# ==========================================
# ASSET-CLASS DETECTION (Industrial & Logistics, Office, ...) — ADDITIVE, NOT
# part of document-type classification above.
#
# This is a second, independent signal layered on top of the RENT_ROLL /
# TRAILING_12 / OFFERING_MEMORANDUM / UNCLASSIFIED routing decision, not a
# replacement or a variant of it. It never changes document_type or scores —
# it answers a different question ("does this content look like a
# specialized commercial asset type — industrial/NNN, office/base-year, etc.
# — as opposed to the residential multifamily shape this pipeline was
# originally built for?") so downstream code (the AI extraction schema in
# views.py) can decide which extra fields to expect at all.
#
# Structured as a ranked multi-class scorer (same pattern as
# CLASSIFICATION_RULES above): every registered asset class gets scored
# independently, and the top scorer wins ONLY if it clears
# MIN_ASSET_CLASS_GROUP_HITS and strictly beats the runner-up — otherwise the
# result is UNDETERMINED. A wrong asset-class guess (or an ambiguous
# Industrial-vs-Office tie) would misdirect the extraction prompt for no
# benefit, so silence (falling back to the standard multifamily-shaped
# extraction) is the safe default. Adding a future asset class (Self-Storage,
# Retail, ...) means adding one new keyword-group constant and one line to
# ASSET_CLASS_RULES below — nothing else in this file changes.
# ==========================================
ASSET_CLASS_INDUSTRIAL = 'INDUSTRIAL'
ASSET_CLASS_OFFICE = 'OFFICE'
ASSET_CLASS_RETAIL = 'RETAIL'
ASSET_CLASS_UNDETERMINED = 'UNDETERMINED'

INDUSTRIAL_KEYWORD_GROUPS = [
    ['square feet', 'square footage', 'sq ft', 'sq. ft.', 'rentable square footage', 'rsf', 'gla'],
    ['nnn', 'triple net', 'triple-net lease', 'net net net'],
    ['cam reimbursement', 'cam recovery', 'common area maintenance', 'cam charges', 'expense reimbursement'],
    ['base rent/sf', 'base rent per sf', 'rent per square foot', 'annual rent per sf', '$/sf'],
    ['escalation', 'rent escalation', 'annual escalation', 'escalation clause'],
    ['warehouse', 'distribution center', 'logistics facility', 'industrial park', 'loading dock', 'dock door', 'clear height'],
]

# Office-specific accounting vocabulary — deliberately distinct from the
# Industrial groups above wherever the real-world terms differ (Base
# Year / Expense Stop / TI Allowance / Loss Factor are office-lease-specific
# concepts with no NNN-industrial equivalent), even though a couple of
# groups (RSF, escalation) legitimately overlap between the two asset
# classes — that overlap is fine; the ranked-scorer picks whichever class
# has the stronger overall signal, and ties fall back to UNDETERMINED.
OFFICE_KEYWORD_GROUPS = [
    ['usable square feet', 'usable sf', 'usf', 'rentable square feet', 'rsf'],
    ['base year', 'base year expenses', 'expense base year'],
    ['expense stop', 'expense stop amount', 'stop amount'],
    ['tenant improvements', 'ti allowance', 'tenant improvement allowance', 'build-out allowance'],
    ['net effective rent', 'effective rent'],
    ['pro-rata share', 'pro rata share', 'proportionate share', 'loss factor'],
    ['escalation clause', 'rent escalation', 'annual escalation'],
]


# Retail-specific accounting vocabulary — Percentage Rent / breakpoint /
# anchor-tenant concepts have no equivalent in the Industrial or Office
# groups above. GLA and CAM legitimately overlap with the other classes
# (same overlap-is-fine reasoning as Office's RSF/escalation groups); the
# ranked scorer resolves it by whichever class has the stronger signal.
RETAIL_KEYWORD_GROUPS = [
    ['gross leasable area', 'gla'],
    ['percentage rent', 'percent rent'],
    ['overage rent', 'natural breakpoint', 'artificial breakpoint', 'breakpoint'],
    ['gross sales threshold', 'tenant sales log', 'gross sales', 'gross annual sales'],
    ['common area maintenance', 'cam charges', 'cam reimbursement', 'cam recovery'],
    ['anchor tenant', 'anchor store', 'shopping center', 'strip center'],
]

ASSET_CLASS_RULES = [
    (ASSET_CLASS_INDUSTRIAL, INDUSTRIAL_KEYWORD_GROUPS),
    (ASSET_CLASS_OFFICE, OFFICE_KEYWORD_GROUPS),
    (ASSET_CLASS_RETAIL, RETAIL_KEYWORD_GROUPS),
]

MIN_ASSET_CLASS_GROUP_HITS = 2


def _detect_asset_class(lower_text: str) -> dict:
    """Additive asset-class signal — see module note above. Never touches document_type."""
    scores = {
        asset_class: _score_text(lower_text, keyword_groups)
        for asset_class, keyword_groups in ASSET_CLASS_RULES
    }

    ranked = sorted(scores.items(), key=lambda kv: kv[1], reverse=True)
    top_class, top_score = ranked[0]
    runner_up_score = ranked[1][1]

    if top_score >= MIN_ASSET_CLASS_GROUP_HITS and top_score > runner_up_score:
        asset_class = top_class
    else:
        asset_class = ASSET_CLASS_UNDETERMINED

    return {"asset_class": asset_class, "scores": scores}


class DocumentClassificationError(Exception):
    """
    Raised when the file can't be read well enough to even attempt
    classification (corrupt file, unreadable encoding, etc.) — distinct from
    a clean UNCLASSIFIED result, which means "read fine, just no confident
    match." Callers should treat this the same as UNCLASSIFIED for routing
    purposes (the user can still place the file manually), just with a more
    specific message.
    """
    pass


def _score_text(lower_text: str, keyword_groups) -> int:
    """Counts how many distinct keyword groups have at least one hit in the given (already-lowercased) text."""
    score = 0
    for group in keyword_groups:
        if any(phrase in lower_text for phrase in group):
            score += 1
    return score


def _sample_pdf(uploaded_file) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(uploaded_file) as pdf:
        for page in pdf.pages[:PDF_PAGES_TO_SCAN]:
            text = page.extract_text()
            if text:
                text_parts.append(text)
            for table in page.extract_tables():
                if table and table[0]:
                    text_parts.append(' '.join(str(cell) for cell in table[0] if cell))
    return '\n'.join(text_parts)


def _sample_docx(uploaded_file) -> str:
    from docx import Document

    doc = Document(uploaded_file)
    paragraphs = [p.text for p in doc.paragraphs[:DOCX_PARAGRAPHS_TO_SCAN] if p.text.strip()]

    table_headers = []
    if doc.tables:
        first_row = doc.tables[0].rows[0] if doc.tables[0].rows else None
        if first_row is not None:
            table_headers = [cell.text for cell in first_row.cells]

    return '\n'.join(paragraphs + table_headers)


def _sample_excel(uploaded_file) -> str:
    import pandas as pd
    from .parsers import detect_header_row, _build_column_names

    header_row_idx = detect_header_row(uploaded_file)

    uploaded_file.seek(0)
    probe = pd.read_excel(
        uploaded_file, sheet_name=0, header=None,
        nrows=header_row_idx + 1 + EXCEL_SAMPLE_ROWS
    )
    uploaded_file.seek(0)

    columns, used_subheader = _build_column_names(probe, header_row_idx)
    data_start = header_row_idx + (2 if used_subheader else 1)
    sample_rows = probe.iloc[data_start:data_start + EXCEL_SAMPLE_ROWS]

    return ' '.join(columns) + '\n' + sample_rows.astype(str).to_string(index=False)


def _sample_csv(uploaded_file) -> str:
    import pandas as pd

    df = pd.read_csv(uploaded_file, nrows=EXCEL_SAMPLE_ROWS)
    uploaded_file.seek(0)
    return ' '.join(str(c) for c in df.columns) + '\n' + df.astype(str).to_string(index=False)


def _extract_classification_sample(uploaded_file) -> str:
    """
    Pulls a small, cheap text sample to classify from. This is intentionally
    NOT the full-document extraction used later for real analysis
    (extract_document_text / extract_t12_text / extract_data_from_excel in
    parsers.py) — classification only needs enough to spot a handful of
    keywords, and re-using the heavier full-document parsers here would slow
    down what's meant to be an instant per-file sort.
    """
    ext = os.path.splitext(uploaded_file.name)[1].lower()
    uploaded_file.seek(0)

    try:
        if ext == '.pdf':
            return _sample_pdf(uploaded_file)
        elif ext in ('.doc', '.docx'):
            return _sample_docx(uploaded_file)
        elif ext in ('.xls', '.xlsx'):
            return _sample_excel(uploaded_file)
        elif ext == '.csv':
            return _sample_csv(uploaded_file)
        else:
            raise DocumentClassificationError(f"Unsupported file type '{ext}' for classification.")
    except DocumentClassificationError:
        raise
    except Exception as e:
        raise DocumentClassificationError(f"Could not read '{uploaded_file.name}' for classification: {str(e)}")
    finally:
        uploaded_file.seek(0)


def classify_document(uploaded_file) -> dict:
    """
    Runs the keyword scan and returns:
        {
            "document_type": one of DOCUMENT_TYPE_RENT_ROLL /
                DOCUMENT_TYPE_TRAILING_12 / DOCUMENT_TYPE_OFFERING_MEMORANDUM /
                DOCUMENT_TYPE_UNCLASSIFIED,
            "scores": {"RENT_ROLL": n, "TRAILING_12": n, "OFFERING_MEMORANDUM": n},
            "asset_class": one of ASSET_CLASS_INDUSTRIAL / ASSET_CLASS_OFFICE /
                ASSET_CLASS_RETAIL / ASSET_CLASS_UNDETERMINED — an ADDITIVE,
                independent signal (see module note above); does not affect
                document_type or scores in any way.
            "asset_class_scores": {"INDUSTRIAL": n, "OFFICE": n, "RETAIL": n},
        }

    Raises DocumentClassificationError if the file itself can't be read.
    """
    sample_text = _extract_classification_sample(uploaded_file)
    lower_text = sample_text.lower()

    scores = {
        doc_type: _score_text(lower_text, keyword_groups)
        for doc_type, keyword_groups in CLASSIFICATION_RULES
    }

    ranked = sorted(scores.items(), key=lambda kv: kv[1], reverse=True)
    top_type, top_score = ranked[0]
    runner_up_score = ranked[1][1]

    if top_score >= MIN_DISTINCT_GROUP_HITS and top_score > runner_up_score:
        document_type = top_type
    else:
        document_type = DOCUMENT_TYPE_UNCLASSIFIED

    asset_class_result = _detect_asset_class(lower_text)

    return {
        "document_type": document_type,
        "scores": scores,
        "asset_class": asset_class_result["asset_class"],
        "asset_class_scores": asset_class_result["scores"],
    }
