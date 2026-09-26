# backend/ai_service/excel_ppm_generator.py
"""
Institutional PPM Master Grid spreadsheet generator.

Deliberately a SEPARATE, isolated file from excel_generator.py rather than an
extension of it — the CRE generator builds a full underwriting workbook with
debt-service formulas, cap-rate calculations, and rent-roll-specific
structure that has no PPM equivalent. This generator does something
different in kind: given several PPM AnalysisReport rows, it builds ONE
comparison grid with one row per analyzed deal, for side-by-side
institutional portfolio review — not a per-deal underwriting model.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Column order for the grid. Each tuple is (metrics_key, header_label).
GRID_COLUMNS = [
    ("property_name", "Asset Name"),
    ("sponsor_platform_name", "Sponsor Platform"),
    ("offering_volume", "Offering Volume"),
    ("loan_to_cost_ratio", "Loan-to-Cost Ratio"),
    ("weighted_exit_cap", "Weighted Exit Cap"),
    ("total_sponsor_fees", "Total Sponsor Fees"),
    ("year_1_yield", "Year 1 Yield"),
    ("reconciliation_flag", "Reconciliation Flag"),
    ("reconciliation_delta", "Sources & Uses Delta ($)"),
    ("created_at", "Analyzed On"),
]

# Header styling matches excel_generator.py's own header row (dark fill,
# bold white text) so a PPM export doesn't look like a different product.
HEADER_FILL = PatternFill(start_color="FF333333", end_color="FF333333", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")

# Reconciliation flag row-coloring — RED reuses the exact fill
# excel_generator.py already uses for flagged reconciliation rows, so a real
# mismatch reads the same way across both products. GREEN/NEEDS_REVIEW are
# new but follow the same soft-pastel convention.
FLAG_FILLS = {
    "GREEN": PatternFill(start_color="FFD4EDDA", end_color="FFD4EDDA", fill_type="solid"),
    "RED": PatternFill(start_color="FFF8D7DA", end_color="FFF8D7DA", fill_type="solid"),
    "NEEDS_REVIEW": PatternFill(start_color="FFFFF3CD", end_color="FFFFF3CD", fill_type="solid"),
}


class PPMExcelGenerationError(Exception):
    """Raised when the Master Grid can't be built from the given reports."""
    pass


def _row_values_for_report(report) -> dict:
    """
    Extracts the flat set of grid values for one AnalysisReport row. Reads
    most fields from report.metrics (the PPM metrics JSON — see
    PPMUnderwriteView for its shape), but property_name and created_at come
    straight off the model since those are real columns, not nested in the
    JSON blob.
    """
    metrics = report.metrics or {}
    reconciliation = metrics.get('reconciliation') or {}

    return {
        "property_name": report.property_name or "Untitled Asset",
        "sponsor_platform_name": metrics.get('sponsor_platform_name'),
        "offering_volume": metrics.get('offering_volume'),
        "loan_to_cost_ratio": metrics.get('loan_to_cost_ratio'),
        "weighted_exit_cap": metrics.get('weighted_exit_cap'),
        "total_sponsor_fees": metrics.get('total_sponsor_fees'),
        "year_1_yield": metrics.get('year_1_yield'),
        "reconciliation_flag": reconciliation.get('flag'),
        "reconciliation_delta": reconciliation.get('delta'),
        "created_at": report.created_at.strftime("%Y-%m-%d") if report.created_at else "",
    }


def build_ppm_master_grid(reports) -> bytes:
    """
    Builds a single-sheet Master Grid workbook, one row per PPM
    AnalysisReport in `reports`. `reports` is expected to already be scoped
    and filtered by the caller (organization + document_type='PPM' — see
    ExportPPMMasterGridView in views.py); this function does no filtering of
    its own and trusts whatever it's given.

    Raises PPMExcelGenerationError if given no reports — an empty grid isn't
    a useful download, better to tell the user plainly than hand them a
    workbook with just a header row.
    """
    reports = list(reports)
    if not reports:
        raise PPMExcelGenerationError("No PPM analyses found to include in the Master Grid.")

    wb = Workbook()
    ws = wb.active
    ws.title = "PPM Master Grid"

    for col_idx, (_, label) in enumerate(GRID_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for row_idx, report in enumerate(reports, start=2):
        values = _row_values_for_report(report)
        for col_idx, (key, _label) in enumerate(GRID_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values.get(key))
            if key == "reconciliation_flag":
                fill = FLAG_FILLS.get(values.get(key))
                if fill:
                    cell.fill = fill
                    cell.font = Font(bold=True)

    for col_idx, (_key, label) in enumerate(GRID_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(label) + 4, 14)

    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
